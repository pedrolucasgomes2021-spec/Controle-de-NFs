import streamlit as st
import pandas as pd
import os
import pyperclip
from datetime import datetime
from openpyxl import Workbook, load_workbook

# -------------------------
# CONFIGURAÇÕES INICIAIS
# -------------------------
EXCEL_FILE = "banco_dados.xlsx"
SHEET_LANC = "Lançamentos"
SHEET_RESP = "Responsáveis"
SHEET_FORN = "Fornecedores"

# Cria o banco se não existir
def inicializar_banco():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()

        ws1 = wb.active
        ws1.title = SHEET_LANC
        ws1.append([
            "Responsável", "Nota Fiscal", "P.O / Linhas", "Nº MIGO",
            "Referência", "Fornecedor", "Data Lançamento",
            "P.O Guarda-Chuva?", "Comprador(a)"
        ])

        wb.create_sheet(SHEET_RESP)
        wb[SHEET_RESP].append(["Responsável"])

        wb.create_sheet(SHEET_FORN)
        wb[SHEET_FORN].append(["Fornecedor"])

        wb.save(EXCEL_FILE)

def carregar_lista(sheet):
    wb = load_workbook(EXCEL_FILE)
    ws = wb[sheet]
    valores = [cell.value for cell in ws["A"] if cell.value not in (None, sheet)]
    return list(set(valores))

def salvar_lista(sheet, novo_valor):
    wb = load_workbook(EXCEL_FILE)
    ws = wb[sheet]
    valores = [cell.value for cell in ws["A"] if cell.value is not None]
    if novo_valor not in valores:
        ws.append([novo_valor])
        wb.save(EXCEL_FILE)

def salvar_lancamentos(lancamentos):
    wb = load_workbook(EXCEL_FILE)
    ws = wb[SHEET_LANC]
    for lanc in lancamentos:
        ws.append(lanc)
    wb.save(EXCEL_FILE)

def abrir_excel():
    caminho = os.path.abspath(EXCEL_FILE)
    os.startfile(caminho)  # abre no Excel no Windows

# -------------------------
# STREAMLIT CONFIG
# -------------------------
st.set_page_config(page_title="Cadastro NF/PO", page_icon="📘", layout="centered")

st.markdown(
    """
    <style>
    .stApp { background-color: #0d1117; color: #dbeafe; }
    .title {color: #06b6d4; font-size:24px; font-weight:700; text-align:center; margin-bottom:20px;}
    .card {background:#161b22; border-radius:10px; padding:12px; margin-bottom:10px; font-size:14px;}
    </style>
    """, unsafe_allow_html=True
)

st.markdown('<div class="title">📘 Cadastro de Notas / P.O</div>', unsafe_allow_html=True)

inicializar_banco()

responsaveis = carregar_lista(SHEET_RESP)
fornecedores = carregar_lista(SHEET_FORN)

# -------------------------
# CAMPOS DINÂMICOS
# -------------------------
if "num_blocos" not in st.session_state:
    st.session_state.num_blocos = 1

if st.button("➕ Mais um lançamento"):
    st.session_state.num_blocos += 1

lancamentos_temp = []

for i in range(st.session_state.num_blocos):
    with st.container():
        st.markdown(f'<div class="card">', unsafe_allow_html=True)
        st.write(f"### Lançamento {i+1}")

        col1, col2 = st.columns(2)
        with col1:
            novo_resp = st.text_input(f"Cadastrar novo responsável {i+1}", key=f"novo_resp_{i}")
            if st.button(f"Salvar responsável {i+1}", key=f"btn_resp_{i}"):
                if novo_resp.strip():
                    salvar_lista(SHEET_RESP, novo_resp.strip())
                    st.rerun()
            responsavel = st.selectbox("Responsável", options=responsaveis, key=f"resp_{i}")

        with col2:
            novo_forn = st.text_input(f"Cadastrar novo fornecedor {i+1}", key=f"novo_forn_{i}")
            if st.button(f"Salvar fornecedor {i+1}", key=f"btn_forn_{i}"):
                if novo_forn.strip():
                    salvar_lista(SHEET_FORN, novo_forn.strip())
                    st.rerun()
            fornecedor = st.selectbox("Fornecedor", options=fornecedores, key=f"forn_{i}")

        nota = st.text_input("Número da Nota Fiscal", key=f"nota_{i}")
        po = st.text_input("P.O / Linhas da P.O", key=f"po_{i}")
        migo = st.text_input("Número MIGO", key=f"migo_{i}")
        referencia = st.text_input("Referência", key=f"ref_{i}")
        data_lanc = st.date_input("Data do Lançamento", value=datetime.today(), key=f"data_{i}")
        guarda_chuva = st.checkbox("É P.O Guarda-Chuva?", key=f"gc_{i}")

        comprador = ""
        if guarda_chuva:
            comprador = st.text_input("Nome do Comprador(a)", key=f"comp_{i}")

        lancamentos_temp.append([
            responsavel, nota, po, migo, referencia,
            fornecedor, data_lanc.strftime("%d/%m/%Y"),
            "Sim" if guarda_chuva else "Não",
            comprador
        ])
        st.markdown('</div>', unsafe_allow_html=True)

# -------------------------
# BOTÕES DE AÇÃO
# -------------------------
col1, col2, col3 = st.columns(3)
with col1:
    if st.button("💾 Lançar no banco de dados"):
        salvar_lancamentos(lancamentos_temp)
        st.success("Lançamentos gravados no banco com sucesso!")

with col2:
    if st.button("🔄 Resetar"):
        st.session_state.num_blocos = 1
        st.rerun()

with col3:
    if st.button("📂 Abrir banco de dados"):
        abrir_excel()

# -------------------------
# GERAR E-MAIL
# -------------------------
if st.button("📧 Gerar e-mail"):
    corpo = "Prezados, bom dia!\n\nSolicito o lançamento das NFs a seguir:\n\n"
    for lanc in lancamentos_temp:
        corpo += f"- NF: {lanc[1]}\n- PO: {lanc[2]}\n- MIGO: {lanc[3]}\n- REFERÊNCIA: {lanc[4]}\n- FORNECEDOR: {lanc[5]}\n\n"
    if any(l[7] == "Sim" for l in lancamentos_temp):
        for lanc in lancamentos_temp:
            if lanc[7] == "Sim" and lanc[8]:
                corpo += f"({lanc[8]}) Favor, realizar ajustes na P.O. RC já modificada.\n\n"
                break
    corpo += "Desde já, agradeço!"

    st.text_area("Mensagem de e-mail:", corpo, height=250)
    if st.button("📋 Copiar e-mail"):
        pyperclip.copy(corpo)
        st.success("Mensagem copiada para a área de transferência!")
        